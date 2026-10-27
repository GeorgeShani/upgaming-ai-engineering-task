from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from app.documents.parsing import MAX_UPLOAD_BYTES, ParseError, parse_upload
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def pdf_bytes(lines=("Payment router failed.", "Recovery took 12 minutes."), encrypted=False):
    writer = PdfWriter()
    for line in lines:
        page = writer.add_blank_page(612, 792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 72 720 Td ({line}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.add_metadata({"/Title": "Fictional test incident", "/Author": "Fixture author"})
    if encrypted:
        writer.encrypt("fixture-password")
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def docx_bytes(table=False):
    doc = Document()
    doc.core_properties.title = "Fictional test incident"
    doc.core_properties.author = "Fixture author"
    doc.add_paragraph("Payment router failed.")
    if table:
        row = doc.add_table(rows=1, cols=2).rows[0]
        row.cells[0].text = "Service"
        row.cells[1].text = "payment-router"
    doc.add_paragraph("Recovery took 12 minutes.")
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def test_three_formats_share_normalized_identity_and_preserve_locations():
    parsed = [
        parse_upload("incident.pdf", pdf_bytes()),
        parse_upload("incident.docx", docx_bytes()),
        parse_upload("incident.txt", b"Payment router failed.\n\nRecovery took 12 minutes."),
    ]
    assert len({doc.content_sha256 for doc in parsed}) == 1
    assert [loc.number for loc in parsed[0].locations] == [1, 2]
    assert parsed[0].locations[0].kind == "page"
    assert parsed[1].locations[0].kind == "paragraph"
    assert parsed[0].metadata["author"] == "Fixture author"
    assert all(doc.text[loc.start : loc.end] == loc.text for doc in parsed for loc in doc.locations)


def test_docx_table_rows_are_extracted_between_paragraphs():
    parsed = parse_upload("incident.docx", docx_bytes(table=True))
    assert [loc.kind for loc in parsed.locations] == ["paragraph", "table_row", "paragraph"]
    assert parsed.locations[1].text == "Service | payment-router"


@pytest.mark.parametrize(
    "name,data,status",
    [
        ("empty.txt", b"", 422),
        ("binary.txt", b"abc\x00def", 422),
        ("encoding.txt", b"\xff\xfeabc", 422),
        ("wrong.pdf", b"not a PDF", 422),
        ("wrong.docx", b"not a ZIP", 422),
        ("script.html", b"<p>text</p>", 415),
        ("large.txt", b"x" * (MAX_UPLOAD_BYTES + 1), 413),
    ],
    ids=["empty", "binary", "encoding", "pdf-signature", "docx-signature", "extension", "oversize"],
)
def test_invalid_uploads_are_rejected(name, data, status):
    with pytest.raises(ParseError) as error:
        parse_upload(name, data)
    assert error.value.status == status


def test_encrypted_and_textless_pdf_are_explicitly_rejected():
    with pytest.raises(ParseError, match="Encrypted"):
        parse_upload("locked.pdf", pdf_bytes(encrypted=True))
    with pytest.raises(ParseError, match="No extractable text"):
        parse_upload("scan.pdf", pdf_bytes(lines=("",)))


def test_docx_expansion_is_bounded_before_document_loading():
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * 20_000_001)
    with pytest.raises(ParseError, match="expanded content"):
        parse_upload("large.docx", buffer.getvalue())


def test_filename_is_display_metadata_not_a_filesystem_path():
    parsed = parse_upload("..\\folder/incident.TXT", b"Safe text.")
    assert parsed.filename == "incident.TXT"
    assert parsed.format == "txt"


def test_utf8_bom_and_non_latin_text_are_supported():
    parsed = parse_upload("incident.txt", "\ufeffქართული ტექსტი.\n\nSecond paragraph.".encode())
    assert parsed.text == "ქართული ტექსტი. Second paragraph."
    assert [loc.number for loc in parsed.locations] == [1, 2]
